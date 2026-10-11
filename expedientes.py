"""Validación y consulta de expedientes. No requiere servicios externos."""

import csv
from dataclasses import dataclass
from datetime import date
from io import StringIO
import unicodedata

REQUIRED_COLUMNS = {"numero", "asunto", "estado", "fecha_ingreso", "fecha_vencimiento"}
STATES = ("pendiente", "en proceso", "observado", "finalizado")


def normalize(value: str) -> str:
    return "".join(
        char for char in unicodedata.normalize("NFD", value.strip().casefold())
        if unicodedata.category(char) != "Mn"
    )


@dataclass(frozen=True)
class Expediente:
    numero: str
    asunto: str
    estado: str
    fecha_ingreso: date
    fecha_vencimiento: date

    def vencido(self, today: date) -> bool:
        return self.estado != "finalizado" and self.fecha_vencimiento < today


@dataclass(frozen=True)
class Incidencia:
    fila: int
    numero: str
    detalle: str


@dataclass
class Importacion:
    expedientes: list[Expediente]
    incidencias: list[Incidencia]
    total_filas: int


def validate_rows(headers, rows) -> Importacion:
    """La fila 1 es el encabezado; los errores conservan la fila de origen."""
    columns = [normalize(str(header)) if header is not None else "" for header in headers]
    if len(columns) != len(set(columns)):
        raise ValueError("El archivo contiene encabezados repetidos.")
    missing = REQUIRED_COLUMNS - set(columns)
    if missing:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(sorted(missing)))
    accepted, issues, seen = [], [], set()
    total = 0
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        total += 1
        if len(values) != len(columns):
            issues.append(Incidencia(row_number, "", "La cantidad de celdas no coincide con el encabezado."))
            continue
        record = {key: str(value).strip() if value is not None else ""
                  for key, value in zip(columns, values)}
        number = record["numero"]
        errors = []
        empty = sorted(key for key in REQUIRED_COLUMNS if not record[key])
        if empty:
            errors.append("Campos vacíos: " + ", ".join(empty))
        state = normalize(record["estado"])
        if state and state not in STATES:
            errors.append("Estado no válido: " + record["estado"])
        dates = {}
        for field in ("fecha_ingreso", "fecha_vencimiento"):
            if record[field]:
                try:
                    if len(record[field]) != 10:
                        raise ValueError
                    dates[field] = date.fromisoformat(record[field])
                except ValueError:
                    errors.append(f"{field}: usa una fecha válida con formato AAAA-MM-DD")
        if len(dates) == 2 and dates["fecha_vencimiento"] < dates["fecha_ingreso"]:
            errors.append("El vencimiento es anterior al ingreso")
        identity = normalize(number)
        if identity and identity in seen:
            errors.append("Número de expediente duplicado")
        if errors:
            issues.append(Incidencia(row_number, number, "; ".join(errors)))
            continue
        seen.add(identity)
        accepted.append(Expediente(number, record["asunto"], state,
                                   dates["fecha_ingreso"], dates["fecha_vencimiento"]))
    return Importacion(accepted, issues, total)


def read_csv(content: bytes) -> Importacion:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("Guarda el CSV con codificación UTF-8.") from error
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(StringIO(text), dialect, strict=True)
    try:
        headers = next(rows)
        return validate_rows(headers, rows)
    except StopIteration as error:
        raise ValueError("El archivo está vacío.") from error
    except csv.Error as error:
        raise ValueError("El CSV tiene un formato incorrecto.") from error


def read_excel(content: bytes) -> Importacion:
    from datetime import datetime
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            raise ValueError("La primera hoja está vacía.")
        def converted_rows():
            for row in rows:
                yield [value.date().isoformat() if isinstance(value, datetime)
                       else value.isoformat() if isinstance(value, date)
                       else value for value in row]
        return validate_rows(headers, converted_rows())
    finally:
        workbook.close()


def filter_records(records, query="", state="Todos", due="Todos", today=None):
    today = today or date.today()
    query = normalize(query)
    result = []
    for record in records:
        if query and query not in normalize(record.numero + " " + record.asunto):
            continue
        if state != "Todos" and record.estado != state:
            continue
        days = (record.fecha_vencimiento - today).days
        if due == "Vencidos" and not record.vencido(today):
            continue
        if due == "Próximos 7 días" and (record.estado == "finalizado" or not 0 <= days <= 7):
            continue
        result.append(record)
    return sorted(result, key=lambda record: (record.fecha_vencimiento, record.numero))


def export_csv(records) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["numero", "asunto", "estado", "fecha_ingreso", "fecha_vencimiento"])
    for record in records:
        # Evitar que Excel interprete texto importado como una fórmula.
        def safe(value):
            text = str(value)
            return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text
        writer.writerow([safe(record.numero), safe(record.asunto), record.estado,
                         record.fecha_ingreso.isoformat(), record.fecha_vencimiento.isoformat()])
    return output.getvalue().encode("utf-8-sig")
